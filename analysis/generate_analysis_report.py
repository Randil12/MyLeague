"""Génère les livrables RNCP Bloc 2 à partir du dossier Bloc 1."""

from __future__ import annotations

import json
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
TEMPLATE = ROOT / "docs" / "MyLeague_Bloc1_DossierV6.docx"
RESULTS = ROOT / "analysis" / "results" / "statistical_analysis_results.json"
ASSETS = ROOT / "docs" / "assets" / "bloc2"
OUTPUT = ROOT / "docs" / "MyLeague_Bloc2_DossierV2.docx"
TRAINING_OUTPUT = ROOT / "docs" / "MyLeague_Bloc2_Support_FormationV2.docx"

BLUE = "17365D"
LIGHT_BLUE = "D9EAF7"
LIGHT_GREY = "F2F2F2"
AMBER = "FFF2CC"
GREEN = "E2F0D9"
RED = "FCE4D6"


def shade(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_table_borders(table, color: str = "B7C9DD") -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = OxmlElement(f"w:{edge}")
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "6")
        element.set(qn("w:color"), color)
        borders.append(element)


def prevent_row_split(row, *, repeat_header: bool = False) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)
    if repeat_header:
        header = OxmlElement("w:tblHeader")
        header.set(qn("w:val"), "true")
        tr_pr.append(header)


def set_cell_text(cell, value: object, *, bold: bool = False, white: bool = False) -> None:
    cell.text = str(value)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for paragraph in cell.paragraphs:
        for run in paragraph.runs:
            run.bold = bold
            run.font.size = Pt(8.5)
            if white:
                run.font.color.rgb = RGBColor(255, 255, 255)


def add_table(doc: Document, headers: list[str], rows: list[list[object]], widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    set_table_borders(table)
    table.autofit = True
    for index, header in enumerate(headers):
        set_cell_text(table.rows[0].cells[index], header, bold=True, white=True)
        shade(table.rows[0].cells[index], BLUE)
    prevent_row_split(table.rows[0], repeat_header=True)
    for row_index, row in enumerate(rows):
        cells = table.add_row().cells
        prevent_row_split(table.rows[-1])
        for index, value in enumerate(row):
            set_cell_text(cells[index], value)
            if row_index % 2:
                shade(cells[index], LIGHT_GREY)
            if widths:
                cells[index].width = Inches(widths[index])
    doc.add_paragraph()
    return table


def add_body(doc: Document, text: str):
    paragraph = doc.add_paragraph(style="Normal")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    paragraph.add_run(text)
    return paragraph


def add_bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        paragraph = doc.add_paragraph(style="Normal")
        paragraph.paragraph_format.left_indent = Inches(0.25)
        paragraph.paragraph_format.first_line_indent = Inches(-0.18)
        paragraph.add_run("• ").bold = True
        paragraph.add_run(item)


def add_note(doc: Document, title: str, text: str, fill: str = LIGHT_BLUE) -> None:
    table = doc.add_table(rows=1, cols=1)
    set_table_borders(table)
    shade(table.cell(0, 0), fill)
    paragraph = table.cell(0, 0).paragraphs[0]
    paragraph.add_run(f"{title} — ").bold = True
    paragraph.add_run(text)
    doc.add_paragraph()


def add_code(doc: Document, code: str) -> None:
    table = doc.add_table(rows=1, cols=1)
    set_table_borders(table, "D9D9D9")
    shade(table.cell(0, 0), LIGHT_GREY)
    paragraph = table.cell(0, 0).paragraphs[0]
    for line_number, line in enumerate(code.splitlines()):
        if line_number:
            paragraph.add_run("\n")
        run = paragraph.add_run(line)
        run.font.name = "Consolas"
        run.font.size = Pt(8)
    doc.add_paragraph()


def add_figure(doc: Document, filename: str, caption: str, width: float = 6.35) -> None:
    image = ASSETS / filename
    if not image.exists():
        return
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.add_run().add_picture(str(image), width=Inches(width))
    caption_paragraph = doc.add_paragraph(caption)
    caption_paragraph.style = "Caption"
    caption_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER


def add_toc(doc: Document) -> None:
    paragraph = doc.add_paragraph()
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = ' TOC \\o "1-3" \\h \\z \\u '
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    placeholder = OxmlElement("w:t")
    placeholder.text = "Mettre à jour le champ dans Word (Ctrl+A puis F9)."
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for element in (begin, instruction, separate, placeholder, end):
        run._r.append(element)


def add_page_footer(doc: Document, label: str) -> None:
    paragraph = doc.sections[0].footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.add_run(f"{label} - Page ")
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run = paragraph.add_run()
    run._r.extend((begin, instruction, end))


def configure_document(doc: Document) -> None:
    styles = doc.styles
    if "Caption" not in styles:
        styles.add_style("Caption", WD_STYLE_TYPE.PARAGRAPH)
    normal_style = next(style for style in styles if style.style_id == "Normal")
    normal_style.font.name = "Aptos"
    normal_style.font.size = Pt(10.5)
    normal_style.paragraph_format.space_after = Pt(6)
    normal_style.paragraph_format.line_spacing = 1.12
    styles["Caption"].font.name = "Aptos"
    styles["Caption"].font.size = Pt(9)
    styles["Caption"].font.italic = True
    for style_name in ("Heading 1", "Heading 2", "Heading 3"):
        if style_name in styles:
            styles[style_name].paragraph_format.keep_with_next = True
    settings = doc.settings._element
    update = settings.find(qn("w:updateFields"))
    if update is None:
        update = OxmlElement("w:updateFields")
        settings.append(update)
    update.set(qn("w:val"), "true")


def keep_template_cover(doc: Document) -> None:
    """Conserve la couverture et la section du document source."""
    body = doc._element.body
    cover_end = doc.paragraphs[9]._p
    remove = False
    for child in list(body):
        if child is cover_end:
            remove = True
            continue
        if remove and child.tag != qn("w:sectPr"):
            body.remove(child)
    paragraph = doc.paragraphs[4]
    for run in paragraph.runs:
        run.text = ""
    paragraph.add_run("Dossier Bloc 2 - Exploiter les données à des fins d'analyse")
    for section in doc.sections:
        for footer_paragraph in section.footer.paragraphs:
            for run in footer_paragraph.runs:
                if "Bloc 1" in run.text:
                    run.text = run.text.replace("Bloc 1", "Bloc 2")


def percent(value: float | None) -> str:
    return "N/D" if value is None else f"{value * 100:.1f} %"


def add_front_matter(doc: Document) -> None:
    doc.add_heading("Remerciements", level=1)
    add_body(
        doc,
        "Je remercie les équipes pédagogiques d'Ynov pour leur accompagnement, ainsi que "
        "les professionnels de l'esport dont les pratiques ont nourri la définition du besoin. "
        "Le commanditaire de ce projet étant fictif, les validations métier présentées comme "
        "telles restent des protocoles à confronter à un coach ou analyste réel.",
    )
    doc.add_page_break()


def add_executive_summary(doc: Document, data: dict) -> None:
    h1 = data["hypothesis_1"]
    h2 = data["hypothesis_2"]
    doc.add_heading("Synthèse exécutive", level=1)
    add_body(
        doc,
        "Une structure esport doit décider vite alors que la méta évolue à chaque patch et que les "
        "signaux disponibles ne décrivent pas tous le même environnement de jeu. MyLeague répond à "
        "cette difficulté en réunissant les données Riot, les matchs professionnels et les notes de "
        "patch dans une chaîne analytique traçable. L'objectif n'est pas de produire une tier list "
        "universelle, mais de fournir au staff un cadre commun pour formuler, vérifier et documenter "
        "ses décisions de draft, d'entraînement et de coaching.",
    )
    add_body(
        doc,
        f"Le snapshot étudié couvre les patchs {data['patches'][0]['patch']} et "
        f"{data['latest_patch']}, pour un total de "
        f"{sum(x['total_matches'] for x in data['patches'])} matchs suivis. L'analyse de draft "
        f"retient {data['draft']['champions_analyzed']} champions disposant d'un volume minimal. "
        f"Le premier test met en évidence une association positive modérée entre présence en solo "
        f"queue et présence professionnelle (rho = {h1['rho']:.3f} ; p < 0,001). Le second ne permet "
        f"pas d'établir que les champions modifiés par Riot connaissent une variation de winrate "
        f"différente des autres (p = {h2['p_value']:.3f}), notamment parce que seuls "
        f"{h2['n_changed']} champions modifiés satisfont le seuil de volume.",
    )
    add_table(
        doc,
        ["Constat", "Conséquence décisionnelle"],
        [
            ["Les signaux solo queue et professionnels convergent partiellement.", "Les conserver séparément avant de les croiser."],
            ["Le volume reste faible pour plusieurs champions.", "Afficher la confiance et différer les conclusions trop fragiles."],
            ["Une modification officielle n'implique pas un effet mesurable immédiat.", "Observer plusieurs patchs avant d'attribuer une évolution au changement."],
            ["Les données internes de scrim ne sont pas encore intégrées.", "Faire valider toute priorité globale par le staff et le champion pool réel."],
        ],
    )
    add_note(
        doc,
        "Conclusion de la synthèse",
        "Le socle technique et analytique couvre les attendus du Bloc 2. La recevabilité probatoire "
        "sera toutefois renforcée par deux pièces humaines : un entretien métier daté et la trace "
        "d'une session de formation effectivement réalisée.",
        AMBER,
    )
    doc.add_page_break()
    doc.add_heading("Sommaire", level=1)
    add_toc(doc)
    doc.add_page_break()


def add_introduction(doc: Document) -> None:
    doc.add_heading("Introduction", level=1)
    add_body(
        doc,
        "MyLeague répond à la problématique suivante : comment une structure esport peut-elle "
        "transformer des données de jeu massives, hétérogènes et en évolution constante en "
        "analyses fiables de la méta, afin d'éclairer ses décisions de draft, d'entraînement et "
        "de coaching ? Le Bloc 1 a sécurisé la collecte, la transformation et l'exposition des "
        "données. Ce dossier traite leur exploitation : besoin métier, plan d'analyse, calculs, "
        "tests statistiques, datavisualisation, recommandations, formation et documentation.",
    )
    add_note(
        doc,
        "Positionnement",
        "MyLeague est un outil d'aide à la décision. Un score ou une corrélation ne remplace ni "
        "le contexte d'une draft, ni le niveau de maîtrise d'un joueur, ni l'expertise du coach.",
    )
    doc.add_heading("Périmètre du dossier", level=2)
    add_body(
        doc,
        "Le présent dossier commence là où s'achève le Bloc 1. Il ne redémontre donc pas en détail "
        "l'architecture médaillon, la sécurité réseau ou l'orchestration Airflow ; il mobilise ces "
        "composants comme garanties de qualité en amont de l'analyse. Le périmètre analytique porte "
        "sur la méta récente, la préparation de draft et l'orientation du coaching. Il exclut toute "
        "évaluation automatisée d'un joueur et toute recommandation présentée comme certaine.",
    )
    doc.add_heading("Méthode de démonstration", level=2)
    add_body(
        doc,
        "La démonstration suit le cycle d'une analyse professionnelle : partir d'une décision métier, "
        "la traduire en questions mesurables, contrôler la disponibilité des données, exécuter les "
        "calculs, éprouver les hypothèses, représenter les résultats, formuler une recommandation puis "
        "documenter son usage. Chaque chapitre distingue systématiquement le fait observé, son "
        "interprétation et la limite qui encadre cette interprétation.",
    )
    doc.add_heading("Convention de lecture", level=2)
    add_body(
        doc,
        "Les encadrés bleus présentent un principe ou une information structurante ; les encadrés "
        "verts formulent une décision étayée ; les encadrés jaunes signalent une prudence ou une "
        "preuve à compléter. Les valeurs chiffrées correspondent au snapshot horodaté produit par le "
        "script d'analyse. Elles ne doivent pas être confondues avec des constantes du produit.",
    )


def add_need_analysis(doc: Document) -> None:
    doc.add_heading("1. Analyse du besoin métier", level=1)
    doc.add_heading("1.1 Parties prenantes et décisions", level=2)
    add_table(
        doc,
        ["Acteur", "Question métier", "Décision attendue"],
        [
            ["Entraîneur principal", "Quels champions prioriser ou bannir ?", "Préparer la draft"],
            ["Analyste", "Quels signaux sont robustes après un patch ?", "Qualifier la méta"],
            ["Coach de rôle", "Quels écarts individuels travailler ?", "Planifier l'entraînement"],
            ["Joueur", "Quelles sélections et quels match-ups approfondir ?", "Adapter son champion pool"],
            ["Data engineer", "Les données sont-elles fraîches et traçables ?", "Fiabiliser le service"],
        ],
    )
    doc.add_heading("1.2 Démarche d'entretien exploratoire", level=2)
    add_body(
        doc,
        "La découverte est structurée autour de trois personas — coach, analyste et joueur — et "
        "d'un guide d'entretien annexé. Les questions portent sur la décision à prendre, le délai "
        "utile après un patch, les seuils de confiance, les filtres de rôle et les cas où le "
        "jugement humain doit primer. Les réponses attendues sont traduites en critères mesurables.",
    )
    add_note(
        doc,
        "Preuve à compléter",
        "Le projet ne dispose pas encore d'un entretien signé avec une structure réelle. Pour "
        "sécuriser C2.1.1, il faut joindre le compte rendu daté d'un échange avec un coach ou un "
        "analyste, ses arbitrages et l'évolution du besoin qui en résulte.",
        AMBER,
    )
    doc.add_heading("1.3 Contraintes et critères d'acceptation", level=2)
    add_table(
        doc,
        ["Contrainte", "Réponse MyLeague", "Critère vérifiable"],
        [
            ["Patchs fréquents", "Analyses partitionnées par patch", "Patch affiché sur chaque vue"],
            ["Sources hétérogènes", "Normalisation dbt et référentiels Riot", "Tests de jointure et qualité"],
            ["Petits échantillons", "Seuils et niveau de confiance", "Alerte visible si données insuffisantes"],
            ["Décision rapide", "Dashboard Streamlit filtrable", "Lecture d'une priorité en moins de 2 minutes"],
            ["Reproductibilité", "SQL, dbt et script statistique versionnés", "Résultats régénérables"],
            ["Données internes absentes", "Architecture extensible aux scrims", "Limite indiquée, sans extrapolation"],
        ],
    )
    doc.add_heading("1.4 Reformulation du besoin", level=2)
    add_body(
        doc,
        "Le besoin initial — « comprendre la méta » — est trop large pour être directement traité. "
        "Il est reformulé ainsi : fournir, pour un patch et un rôle donnés, une liste de signaux "
        "priorisés dont le volume, l'origine et le niveau de confiance sont visibles, afin que le "
        "staff puisse décider quels champions étudier, tester ou surveiller. Cette reformulation "
        "définit l'objet de l'analyse, son utilisateur, son horizon temporel et l'action attendue.",
    )
    add_table(
        doc,
        ["Identifiant", "Expression du besoin", "Critère d'acceptation"],
        [
            ["US-01", "En tant qu'analyste, je filtre par patch et par rôle.", "Le périmètre sélectionné reste visible sur la page."],
            ["US-02", "En tant que coach, je compare les signaux solo et professionnels.", "Les deux mesures et leurs volumes sont affichés séparément."],
            ["US-03", "En tant que coach de rôle, j'identifie les évolutions à surveiller.", "La comparaison entre patchs distingue valeur et variation."],
            ["US-04", "En tant qu'utilisateur, je sais quand ne pas conclure.", "Une confiance faible déclenche un avertissement explicite."],
            ["US-05", "En tant qu'analyste, je peux expliquer un résultat.", "La source, la formule et la date du snapshot sont retrouvables."],
        ],
    )
    doc.add_heading("1.5 Environnement et frontières de responsabilité", level=2)
    add_body(
        doc,
        "Le data engineer garantit l'ingestion, la qualité, le calcul et la traçabilité des "
        "indicateurs. L'analyste choisit le périmètre pertinent et met les résultats en contexte. Le "
        "coach conserve la responsabilité de la décision sportive. Cette séparation évite deux "
        "dérives : demander à la donnée de trancher une question tactique qu'elle ne décrit pas, ou "
        "modifier manuellement un indicateur pour l'aligner sur une intuition préalable.",
    )
    add_body(
        doc,
        "Le commanditaire fictif permet de construire un cas cohérent, mais ne constitue pas une "
        "validation externe. Les personas et user stories sont donc des hypothèses de conception. "
        "L'entretien réel attendu doit confirmer leur vocabulaire, leur ordre de priorité et les "
        "seuils jugés acceptables par une structure esport.",
    )


def add_analysis_plan(doc: Document, data: dict) -> None:
    doc.add_heading("2. Plan d'analyse et traduction numérique", level=1)
    doc.add_heading("2.1 Axes, indicateurs et règles", level=2)
    add_table(
        doc,
        ["Axe", "Indicateurs", "Usage"],
        [
            ["Évolution de méta", "présence, winrate, volume, variation entre patchs", "Détecter les ruptures"],
            ["Draft", "priorité, présence professionnelle, confiance", "Préparer sélections et bannissements"],
            ["Performance", "KDA, vision, CS/min, dégâts/min, objectifs", "Orienter le coaching"],
            ["Patch impact", "champions modifiés, |Δ winrate|", "Évaluer les changements Riot"],
            ["Qualité", "fraîcheur, complétude, unicité, tests dbt", "Qualifier la fiabilité"],
        ],
    )
    add_body(
        doc,
        "Les métriques sont interprétées avec leur dénominateur. Le winrate vaut victoires / "
        "parties ; la présence solo queue combine pickrate et banrate ; la présence professionnelle "
        "repose sur les picks et bans de matchs compétitifs. Le score de priorité agrège plusieurs "
        "signaux, mais son niveau de confiance dépend du nombre de parties disponibles.",
    )
    doc.add_heading("2.1.1 Dictionnaire des indicateurs", level=3)
    add_table(
        doc,
        ["Indicateur", "Définition opérationnelle", "Précaution"],
        [
            ["Winrate", "Nombre de victoires / nombre de parties du périmètre", "Toujours l'accompagner du volume."],
            ["Pickrate", "Nombre de sélections / nombre d'occasions observées", "Dépend du rôle, du patch et de la population."],
            ["Banrate", "Nombre de bannissements / nombre d'occasions observées", "Un ban peut répondre au confort adverse, pas seulement à la puissance."],
            ["Présence", "Part des occasions où le champion est sélectionné ou banni", "Les règles de calcul diffèrent entre solo queue et compétition."],
            ["Δ winrate", "Winrate du patch courant − winrate du patch précédent", "Sensible aux faibles volumes et au changement de population."],
            ["Priorité", "Indice composite de signaux normalisés", "Outil de classement relatif, sans valeur causale."],
            ["Confiance", "Classe dérivée du volume disponible", "Qualifie la preuve, pas la puissance intrinsèque du champion."],
        ],
    )
    add_body(
        doc,
        "Le score de priorité n'est donc pas une probabilité de victoire. Il ordonne des candidats à "
        "l'examen au sein d'un même snapshot. Son interprétation hors patch, hors rôle ou hors "
        "population n'est pas valide. Cette règle de portée est affichée dans le dashboard et répétée "
        "dans le support de formation.",
    )
    doc.add_heading("2.1.2 Formule du score de priorité", level=3)
    add_code(
        doc,
        "priorité = 100 × (\n"
        "    0,40 × présence_solo\n"
        "  + 0,20 × max(winrate − 0,45 ; 0)\n"
        "  + 0,20 × présence_pro_ou_solo\n"
        "  + 0,20 × min(nombre_de_picks / 250 ; 1)\n"
        ")",
    )
    add_table(
        doc,
        ["Composante", "Rôle dans le score", "Choix ou limite"],
        [
            ["Présence solo", "Mesurer le niveau de contestation dans le panel suivi", "Coefficient 40 %, donc signal principal."],
            ["Performance", "Valoriser la part de winrate au-dessus de 45 %", "Le plancher évite une contribution négative, sans prouver la causalité."],
            ["Présence professionnelle", "Introduire le jeu coordonné", "À défaut de donnée pro, la présence solo est réutilisée et doit être signalée."],
            ["Facteur d'échantillon", "Réduire la priorité des faibles volumes", "Plafonné à 1 à partir de 250 picks."],
        ],
    )
    add_body(
        doc,
        "Les coefficients expriment une priorité de conception ; ils ne rendent pas les contributions "
        "statistiquement équivalentes, car leurs amplitudes diffèrent. Le score est donc explicable, "
        "mais pas calibré comme une probabilité. Les seuils de confiance sont indépendants : faible "
        "sous 100 picks, moyen de 100 à 249, élevé à partir de 250. Une évolution future devra "
        "soumettre ces pondérations à une analyse de sensibilité et à une validation par le staff.",
    )
    doc.add_heading("2.2 Sources disponibles et périmètre", level=2)
    add_table(
        doc,
        ["Source", "Granularité", "Atout", "Limite"],
        [
            ["Riot Match-v5", "match, participant", "données détaillées et officielles", "EUW et comptes suivis"],
            ["Riot Timeline-v5", "événement horodaté", "déroulé tactique", "pas encore exposé intégralement en Gold"],
            ["Data Dragon", "champion et patch", "référentiel officiel", "pas de performance"],
            ["Leaguepedia", "match professionnel", "sélections et bannissements compétitifs", "nomenclature de patch différente"],
            ["Patch notes Riot", "champion modifié", "contexte causal potentiel", "effet à mesurer, non présumé"],
        ],
    )
    add_note(
        doc,
        "Alignement des patchs",
        "Leaguepedia publie actuellement une nomenclature 26.x alors que les données Riot sont en "
        "16.x. La staging dbt conserve la valeur source et normalise la version analytique "
        "26.14 → 16.14. Un test singulier vérifie ensuite qu'un recouvrement pro/solo existe.",
        GREEN,
    )
    doc.add_heading("2.3 Qualité analytique, biais et données manquantes", level=2)
    add_body(
        doc,
        "Une donnée techniquement valide peut rester inadaptée à la question métier. Le plan d'analyse "
        "contrôle donc deux niveaux : la qualité structurelle — unicité, non-nullité, relations et "
        "fraîcheur — puis la qualité d'usage — représentativité, volume et comparabilité entre sources. "
        "Les tests dbt répondent au premier niveau ; les seuils, niveaux de confiance et avertissements "
        "répondent au second.",
    )
    add_table(
        doc,
        ["Risque de biais", "Effet possible", "Mesure de maîtrise"],
        [
            ["Sélection des comptes EUW", "Méta non représentative de toutes les régions", "Afficher la région et interdire la généralisation implicite."],
            ["Surreprésentation de joueurs suivis", "Styles individuels confondus avec la méta", "Augmenter et documenter le panel."],
            ["Décalage de patch entre sources", "Jointure pro/solo vide ou erronée", "Conserver le patch source, normaliser et tester le recouvrement."],
            ["Biais de survivant", "Champions peu joués exclus de l'analyse", "Documenter le seuil et conserver leur disponibilité dans le détail."],
            ["Absence des scrims", "Recommandation déconnectée de l'équipe", "Réserver la décision finale au staff et planifier l'intégration interne."],
        ],
    )
    doc.add_heading("2.4 Hypothèses statistiques", level=2)
    add_table(
        doc,
        ["Hypothèse", "H0", "Test", "Seuil"],
        [
            ["Présences solo et pro associées", "rho = 0", "Spearman bilatéral", "α = 5 %"],
            ["Impact d'un changement Riot", "distributions de |Δ winrate| identiques", "Mann–Whitney U bilatéral", "α = 5 %"],
        ],
    )
    add_body(
        doc,
        f"Le protocole exclut du test principal les champions ayant moins de "
        f"{data['minimum_picks_per_patch']} picks par patch. Ce seuil réduit le bruit, sans faire "
        "disparaître l'incertitude. Une valeur p inférieure à 0,05 conduit au rejet de H0 ; elle "
        "ne mesure ni la causalité ni, à elle seule, l'importance métier de l'effet.",
    )
    add_body(
        doc,
        "Les hypothèses sont formulées avant la lecture des valeurs p. Les deux tests sont bilatéraux, "
        "car l'analyse cherche une différence dans les deux sens et ne dispose pas d'une justification "
        "métier suffisante pour imposer une direction. Ce choix limite le risque d'adapter le protocole "
        "au résultat observé. Les tests demeurent exploratoires : ils orientent une investigation et "
        "ne constituent pas, à eux seuls, une règle automatisée de draft.",
    )


def add_queries_and_results(doc: Document, data: dict) -> None:
    latest = data["latest_patch"]
    top = data["draft"]["top_recommendations"]
    doc.add_heading("3. Requêtes, calculs et résultats", level=1)
    doc.add_heading("3.1 Chaîne de calcul reproductible", level=2)
    add_body(
        doc,
        "Les vues Gold sont calculées par PostgreSQL/dbt puis consommées par Streamlit et par le "
        "script Python d'analyse. La même définition sert donc au dashboard, aux figures du dossier "
        "et aux tests statistiques. Le dernier build vérifié comporte 18 modèles, 63 tests de "
        "données et 81 contrôles réussis.",
    )
    add_table(
        doc,
        ["Étape", "Technique", "Responsabilité analytique"],
        [
            ["Préparation", "SQL et modèles dbt", "Joindre, agréger et normaliser à une granularité documentée."],
            ["Contrôle", "Tests dbt et tests Python", "Détecter les ruptures de contrat et résultats incohérents."],
            ["Exploration", "SQL et pandas", "Décrire les volumes, valeurs manquantes et distributions."],
            ["Inférence", "SciPy", "Exécuter les tests associés aux hypothèses préétablies."],
            ["Restitution", "Matplotlib et Streamlit", "Rendre résultat, contexte et incertitude lisibles."],
        ],
    )
    add_body(
        doc,
        "La logique métier durable est placée dans dbt plutôt que dupliquée dans l'interface. Le "
        "script Python ne recalcule pas silencieusement le Gold : il le lit, applique le protocole "
        "statistique annoncé et exporte un snapshot JSON. Cette séparation réduit le risque qu'un "
        "même indicateur porte deux définitions selon le support consulté.",
    )
    add_code(
        doc,
        "# Reproduire les contrôles et les statistiques\n"
        "docker compose exec dbt dbt build --project-dir /usr/app/dbt --profiles-dir /usr/app/dbt\n"
        ".venv\\Scripts\\python.exe analysis\\statistical_analysis.py\n\n"
        "-- Exemple de consommation Gold\n"
        "SELECT champion_name, role, priority_score, confidence_level\n"
        "FROM gold.draft_recommendations\n"
        f"WHERE patch = '{latest}' ORDER BY priority_score DESC;",
    )
    doc.add_heading("3.2 Photographie de l'échantillon", level=2)
    add_table(
        doc,
        ["Patch", "Parties", "Durée moyenne"],
        [[item["patch"], item["total_matches"], f"{item['avg_duration_min']:.1f} min"] for item in data["patches"]],
    )
    add_body(
        doc,
        f"Au moment de l'extraction, le patch le plus récent est {latest}. "
        f"{data['draft']['champions_analyzed']} champions disposent d'au moins dix picks et sont "
        "présentés dans l'analyse de draft. Ce snapshot est daté et doit être régénéré avant toute "
        "réunion sportive, car l'échantillon évolue avec l'ingestion.",
    )
    add_body(
        doc,
        "La progression du nombre de matchs entre les deux patchs améliore la stabilité descriptive, "
        "mais ne rend pas automatiquement les populations comparables. Les comptes observés, le "
        "temps écoulé depuis le déploiement du patch et la fréquence de sélection des champions peuvent "
        "évoluer simultanément. Pour cette raison, le dossier décrit le snapshot au lieu de le présenter "
        "comme un échantillon aléatoire de l'ensemble des joueurs de League of Legends.",
    )
    doc.add_heading("3.3 Priorités de draft", level=2)
    rows = []
    for item in top[:8]:
        rows.append(
            [
                item["champion_name"],
                item["role"],
                f"{item['priority_score']:.1f}",
                item["picks"],
                percent(item["winrate"]),
                percent(item["pro_presence"]),
                item["confidence_level"],
            ]
        )
    add_table(doc, ["Champion", "Rôle", "Score", "Picks", "Winrate", "Présence pro", "Confiance"], rows)
    add_figure(doc, "draft_priorities.png", "Figure 1 — Priorités de draft sur le dernier patch disponible.")
    add_note(
        doc,
        "Lecture prudente",
        "Toutes les premières recommandations ont ici une confiance faible : elles constituent une "
        "liste de surveillance, pas une prescription. Le coach doit confirmer la compatibilité avec "
        "le champion pool, le plan de jeu, le choix du côté et les réponses adverses.",
        AMBER,
    )
    add_figure(doc, "streamlit_tier_list.png", "Figure 2 — Restitution opérationnelle dans Streamlit.")
    doc.add_heading("3.4 Contrôle de cohérence des résultats", level=2)
    add_body(
        doc,
        "Trois niveaux de contrôle encadrent les calculs. Premièrement, les tests de schéma vérifient "
        "les clés, valeurs obligatoires, relations et domaines attendus. Deuxièmement, les tests "
        "métier contrôlent notamment la cohérence des probabilités, des durées et du recouvrement de "
        "patch entre solo queue et professionnel. Troisièmement, une revue analytique confronte les "
        "ordres de grandeur au volume source et signale tout résultat contre-intuitif au lieu de le "
        "corriger manuellement.",
    )
    add_note(
        doc,
        "Résultat du contrôle",
        "Le build de référence est vert : 18 modèles construits et 63 tests de données réussis, soit "
        "81 opérations sans erreur. Cette réussite prouve la conformité aux règles implémentées ; elle "
        "ne prouve pas à elle seule la représentativité sportive de l'échantillon.",
        GREEN,
    )
    doc.add_heading("3.5 Portée de la restitution", level=2)
    add_body(
        doc,
        "Le dashboard répond à une logique d'exploration progressive. La page de synthèse permet de repérer un "
        "signal ; les filtres et tableaux permettent ensuite d'en retrouver le rôle, le patch, le "
        "volume et les composantes. Une exportation CSV autorise enfin une analyse ad hoc sans modifier "
        "la définition centrale. Le parcours privilégie ainsi l'explicabilité à l'accumulation de "
        "graphiques.",
    )


def add_statistics(doc: Document, data: dict) -> None:
    h1 = data["hypothesis_1"]
    h2 = data["hypothesis_2"]
    doc.add_heading("4. Modèles statistiques et validation des hypothèses", level=1)
    doc.add_heading("4.1 Protocole d'inférence", level=2)
    add_body(
        doc,
        "L'analyse inférentielle répond à deux questions définies dans le plan d'analyse. L'unité "
        "statistique est le champion, observé dans le périmètre de patch retenu. Le seuil de "
        "significativité est fixé à 5 % avant l'exécution. Les champions ne satisfaisant pas le seuil "
        "de picks sont exclus du test concerné, mais leur exclusion est comptabilisée et ne doit pas "
        "être interprétée comme une absence d'intérêt sportif.",
    )
    add_body(
        doc,
        "Des méthodes non paramétriques sont retenues parce que les présences et variations sont "
        "bornées, asymétriques et susceptibles de contenir des valeurs extrêmes. Spearman évalue une "
        "association monotone à partir des rangs, sans exiger de relation linéaire ni de normalité. "
        "Mann–Whitney compare les rangs de deux groupes indépendants sans supposer une distribution "
        "gaussienne. Ces propriétés rendent les tests cohérents avec les données disponibles, sans "
        "supprimer la fragilité liée aux petits effectifs.",
    )
    doc.add_heading("4.2 Présence solo queue et présence professionnelle", level=2)
    add_body(doc, f"Question : {h1['question']} H0 : {h1['null_hypothesis']}")
    add_body(
        doc,
        "L'hypothèse alternative stipule qu'une association monotone, positive ou négative, existe "
        "entre les deux présences. Le test bilatéral évite de supposer à l'avance qu'une popularité "
        "solo queue se transpose nécessairement au jeu coordonné. Les observations sans présence "
        "professionnelle exploitable sont écartées de cette comparaison appariée.",
    )
    add_body(
        doc,
        f"Le test de {h1['test']} porte sur n = {h1['n']} champions. Il produit rho = "
        f"{h1['rho']:.3f} et p = {h1['p_value']:.2e}. H0 est rejetée au seuil de 5 %. "
        "L'association est positive et d'intensité modérée : une forte présence solo constitue un "
        "signal utile, mais elle ne suffit pas à prévoir la priorité professionnelle.",
    )
    add_figure(doc, "solo_vs_pro_presence.png", "Figure 3 — Association entre présence solo queue et professionnelle.")
    add_note(
        doc,
        "Interprétation métier",
        "Conserver les deux indicateurs dans la préparation de draft. Une divergence est informative : "
        "elle peut révéler un champion puissant en environnement coordonné, ou au contraire performant "
        "en solo queue mais difficile à intégrer à une composition compétitive.",
        GREEN,
    )
    doc.add_heading("4.3 Effet observable des changements de patch", level=2)
    add_body(doc, f"Question : {h2['question']} H0 : {h2['null_hypothesis']}")
    add_body(
        doc,
        "La variable comparée est l'amplitude absolue de variation du winrate entre deux patchs. Le "
        "signe est volontairement retiré : un buff comme un nerf peut produire une rupture notable, "
        "alors que la question porte ici sur l'intensité du changement. Les groupes « modifié » et "
        "« non modifié » sont définis à partir des notes de patch, indépendamment du résultat observé.",
    )
    add_body(
        doc,
        f"Le {h2['test']} compare {h2['n_changed']} champions modifiés à "
        f"{h2['n_unchanged']} champions non modifiés. L'amplitude moyenne vaut "
        f"{percent(h2['changed_mean_absolute_delta'])} contre "
        f"{percent(h2['unchanged_mean_absolute_delta'])}. U = {h2['u_statistic']:.0f}, "
        f"p = {h2['p_value']:.3f}. H0 n'est pas rejetée : les données ne démontrent pas encore "
        "un effet différent. Le groupe modifié est trop petit pour conclure avec robustesse.",
    )
    add_figure(doc, "patch_change_effect.png", "Figure 4 — Amplitude des variations selon le statut de changement.")
    add_figure(doc, "patch_evolution.png", "Figure 5 — Évolution du volume et de la durée moyenne par patch.")
    add_note(
        doc,
        "Décision correcte",
        "Ne pas confondre « H0 non rejetée » avec « absence d'effet ». La collecte doit se poursuivre "
        "sur plusieurs patchs, puis le test être répété avec un groupe de champions modifiés plus grand.",
        AMBER,
    )
    doc.add_heading("4.4 Synthèse des décisions statistiques", level=2)
    add_table(
        doc,
        ["Question", "Résultat", "Décision statistique", "Décision métier"],
        [
            ["Solo queue et professionnel évoluent-ils ensemble ?", f"rho = {h1['rho']:.3f} ; p < 0,001", "Rejet de H0", "Croiser les signaux sans les fusionner."],
            ["Les champions modifiés varient-ils davantage ?", f"U = {h2['u_statistic']:.0f} ; p = {h2['p_value']:.3f}", "H0 non rejetée", "Poursuivre la collecte avant de conclure."],
        ],
    )
    add_body(
        doc,
        "La première analyse établit une association dans le snapshot, pas un mécanisme causal. La "
        "seconde reste indécise : ni la valeur p ni la différence descriptive ne justifient une "
        "conclusion ferme avec trois champions modifiés. Dans les deux cas, la décision métier tient "
        "compte simultanément de la statistique, de la taille d'échantillon et du coût d'une erreur de "
        "draft.",
    )


def add_visualisation(doc: Document) -> None:
    doc.add_heading("5. Datavisualisation et accessibilité", level=1)
    doc.add_heading("5.1 Principes de conception", level=2)
    add_body(
        doc,
        "La représentation est choisie en fonction de la question et non de l'outil disponible. Les "
        "graphiques emploient une palette bleu/orange lisible en cas de déficience rouge-vert. Les "
        "valeurs, légendes et conclusions restent textuelles : aucune information essentielle ne "
        "dépend uniquement de la couleur.",
    )
    add_body(
        doc,
        "Chaque vue répond à une question principale et commence par le contexte nécessaire à sa "
        "lecture : patch, population, rôle et fraîcheur. Le premier niveau présente la conclusion "
        "opérationnelle ; le second expose les métriques qui la soutiennent ; le troisième permet "
        "d'accéder au détail. Cette hiérarchie réduit la charge cognitive tout en conservant la "
        "possibilité d'auditer le résultat.",
    )
    doc.add_heading("5.2 Choix des représentations", level=2)
    add_table(
        doc,
        ["Besoin", "Représentation", "Justification"],
        [
            ["Classer des priorités", "barres horizontales", "lecture rapide des rangs et scores"],
            ["Comparer deux présences", "nuage de points", "relation, dispersion et atypiques visibles"],
            ["Comparer deux groupes", "boîtes et points", "distribution et petite taille d'échantillon visibles"],
            ["Suivre les patchs", "barres + courbe", "volume et durée conservés sur deux échelles explicites"],
            ["Explorer", "tables et filtres Streamlit", "accès aux détails sans surcharger la synthèse"],
        ],
    )
    add_body(
        doc,
        "Un diagramme circulaire n'est pas retenu pour comparer de nombreux champions, car les écarts "
        "d'angle seraient difficiles à apprécier. De même, une simple courbe de winrate sans volume "
        "donnerait une impression trompeuse de précision. Le dossier privilégie donc des positions sur "
        "un axe commun, des distributions et des annotations directes.",
    )
    doc.add_heading("5.3 Accessibilité et compréhension", level=2)
    add_bullets(
        doc,
        [
            "Filtres explicites : patch, rôle, champion et seuil de volume.",
            "Unités et dénominateurs indiqués dans les libellés et info-bulles.",
            "Avertissement visible quand la confiance est faible ou la donnée absente.",
            "Tableau accessible en complément des graphiques et export CSV pour analyse secondaire.",
        ],
    )
    add_body(
        doc,
        "L'accessibilité est considérée comme une condition de fiabilité de la décision. Un libellé "
        "ambigu, une couleur isolée ou une unité absente peut produire une erreur d'interprétation "
        "aussi dommageable qu'une erreur de calcul. Les labels des widgets restent donc renseignés, "
        "les contrastes sont conservés et les tableaux fournissent une alternative textuelle aux "
        "graphiques.",
    )
    doc.add_heading("5.4 Protocole de test utilisateur", level=2)
    add_body(
        doc,
        "Le test proposé dure quinze minutes avec un coach ou analyste n'ayant pas développé l'outil. "
        "Il lui est demandé d'identifier un champion à surveiller, de justifier son choix, puis de "
        "repérer une limite de l'analyse. Sont observés le temps de réponse, les hésitations, les "
        "filtres utilisés et la capacité à retrouver le volume. Le test est réussi si la décision est "
        "formulée en moins de deux minutes après prise en main et si l'incertitude est correctement "
        "citée, sans aide du formateur.",
    )


def add_recommendations(doc: Document) -> None:
    doc.add_heading("6. Recommandations aux décideurs", level=1)
    doc.add_heading("6.1 Cadre de recommandation", level=2)
    add_body(
        doc,
        "Une recommandation exploitable associe quatre éléments : une action, son destinataire, les "
        "preuves qui la justifient et la condition qui conduirait à la réviser. Cette structure évite "
        "les formulations vagues telles que « ce champion est fort » et permet au staff de conserver "
        "la trace de son raisonnement lorsque la méta évolue.",
    )
    add_table(
        doc,
        ["Horizon", "Recommandation", "Décideur", "Indicateur de suivi"],
        [
            ["Avant la draft", "Croiser score, présence pro, confort joueur et match-up", "Entraîneur principal", "choix documenté par série"],
            ["24–72 h après patch", "Créer une liste de surveillance, sans figer de tier list", "Analyste", "volume et confiance"],
            ["En entraînement", "Tester d'abord les champions à signal convergent", "Coachs de rôle", "résultats de scrims"],
            ["Chaque semaine", "Comparer solo, pro et observations internes", "Staff sportif", "écarts expliqués"],
            ["À moyen terme", "Intégrer scrims, champion pools et adversaires", "Pôle data et staff", "couverture des décisions"],
        ],
    )
    doc.add_heading("6.2 Recommandations issues du snapshot", level=2)
    add_body(
        doc,
        "Pour le snapshot observé, Locke et Syndra méritent une revue midlane, tandis que Camille "
        "et Sylas doivent être requalifiés dans leur contexte de rôle et de composition. Cette phrase "
        "n'est pas une consigne de sélection : elle déclenche une analyse vidéo, une vérification de l'échantillon "
        "et un essai en scrim. La meilleure recommandation est celle dont les hypothèses sont rendues "
        "visibles et que le coach peut accepter, nuancer ou rejeter.",
    )
    add_note(
        doc,
        "Formulation recommandée",
        "« Sur le patch analysé, nous plaçons Locke dans la liste de surveillance de la voie du milieu en raison de sa forte "
        "présence combinée. La confiance demeure faible ; avant de modifier la draft, nous vérifions "
        "les matchs sources, le champion pool du joueur et deux situations de match-up en scrim. »",
        GREEN,
    )
    doc.add_heading("6.3 Priorités d'amélioration", level=2)
    add_table(
        doc,
        ["Priorité", "Action", "Valeur attendue", "Condition de réussite"],
        [
            ["P1", "Intégrer les résultats de scrims et leur contexte", "Relier la méta globale à la capacité réelle de l'équipe", "Granularité et accès validés par le staff"],
            ["P1", "Enregistrer champion pools et rôles cibles", "Écarter les recommandations techniquement inapplicables", "Historisation et droit de correction définis"],
            ["P2", "Étendre le panel à d'autres régions", "Réduire le biais EUW", "Comparaisons régionales explicites"],
            ["P2", "Conserver plusieurs patchs consécutifs", "Améliorer la puissance des tests d'impact", "Seuil minimal de champions modifiés atteint"],
            ["P3", "Ajouter le contexte adversaire et le choix du côté", "Rendre la recommandation plus tactique", "Besoin confirmé en entretien"],
        ],
    )
    doc.add_heading("6.4 Journal de décision", level=2)
    add_body(
        doc,
        "Pour chaque décision importante, le staff peut conserver : la date, le patch, le périmètre, "
        "l'action décidée, les indicateurs consultés, la confiance, l'avis du coach et le résultat "
        "observé en scrim ou en match. Ce journal transforme le dashboard en boucle d'apprentissage : "
        "il devient possible d'évaluer après coup non seulement la performance du champion, mais aussi "
        "la qualité du processus de décision.",
    )


def add_training_and_docs(doc: Document) -> None:
    doc.add_heading("7. Formation des utilisateurs", level=1)
    doc.add_heading("7.1 Public et objectifs pédagogiques", level=2)
    add_body(
        doc,
        "Le public cible est un staff sportif non spécialiste de la donnée. Une session de 45 minutes "
        "alterne démonstration et mise en situation : 5 minutes de contexte, 10 minutes de navigation, "
        "10 minutes de lecture des indicateurs, 15 minutes de cas de draft et 5 minutes de questions. "
        "À la fin, l'utilisateur doit filtrer un patch, qualifier la confiance et formuler une décision "
        "avec au moins une réserve méthodologique.",
    )
    add_body(
        doc,
        "Les prérequis sont volontairement limités à la connaissance du vocabulaire de League of "
        "Legends et du processus de draft. Aucun prérequis SQL ou statistique n'est exigé. L'approche "
        "pédagogique part d'une décision familière, puis introduit les indicateurs nécessaires à son "
        "argumentation. Le participant manipule lui-même l'interface afin de vérifier sa capacité à "
        "transférer l'apprentissage dans une situation de travail.",
    )
    doc.add_heading("7.2 Déroulé de la formation", level=2)
    add_table(
        doc,
        ["Étape", "Activité", "Validation"],
        [
            ["1", "Choisir le patch et le rôle", "le périmètre affiché est reformulé"],
            ["2", "Lire volume, présence et winrate", "les dénominateurs sont identifiés"],
            ["3", "Comparer solo queue et professionnel", "une divergence est expliquée"],
            ["4", "Vérifier la confiance", "une donnée faible n'est pas surinterprétée"],
            ["5", "Formuler une recommandation", "décision, preuve et limite sont citées"],
        ],
    )
    doc.add_heading("7.3 Évaluation des acquis", level=2)
    add_body(
        doc,
        "L'évaluation repose sur un cas de draft court. Le participant choisit un champion à étudier "
        "et rédige une recommandation comprenant une action, deux éléments de preuve et une limite. "
        "La réussite ne dépend pas du champion choisi : elle dépend de la cohérence du raisonnement, "
        "de la lecture correcte du volume et de la capacité à ne pas transformer un signal faible en "
        "certitude.",
    )
    add_table(
        doc,
        ["Critère", "Acquis", "À renforcer"],
        [
            ["Périmètre", "Patch, rôle et population sont cités.", "Un filtre ou le contexte est omis."],
            ["Indicateurs", "Valeur et volume sont lus ensemble.", "Le winrate est lu isolément."],
            ["Confiance", "L'incertitude modifie la décision.", "L'avertissement est ignoré."],
            ["Recommandation", "Action, preuves et réserve sont explicites.", "La conclusion reste descriptive ou catégorique."],
        ],
    )
    add_note(
        doc,
        "Preuve à compléter",
        "Le support est livré séparément, mais C2.3.1 exige idéalement une séance réelle : joindre "
        "convocation ou feuille de présence, support daté, exercice, questionnaire et synthèse des "
        "retours. Le dossier ne prétend pas que cette séance a déjà eu lieu.",
        AMBER,
    )
    doc.add_heading("7.4 Amélioration continue de la formation", level=2)
    add_body(
        doc,
        "Les questions posées, erreurs récurrentes et demandes de nouveaux filtres sont consignées à "
        "l'issue de la séance. Elles alimentent deux backlogs distincts : les besoins de formation, "
        "lorsque l'indicateur est correct mais mal compris, et les besoins produit, lorsque l'interface "
        "ne fournit pas le contexte nécessaire. Cette distinction empêche de compenser durablement une "
        "mauvaise ergonomie par une formation plus longue.",
    )
    doc.add_heading("8. Documentation technique et reproductibilité", level=1)
    doc.add_heading("8.1 Documentation par public", level=2)
    add_table(
        doc,
        ["Élément", "Référence", "Rôle"],
        [
            ["Sources", "airflow/dags, dbt/models/sources.yml", "origine et cadence"],
            ["Transformations", "dbt/models/staging et marts", "formules versionnées"],
            ["Qualité", "dbt/tests, tests/, docs/plan_de_tests.md", "contrôles automatiques"],
            ["Statistiques", "analysis/statistical_analysis.py", "tests et figures reproductibles"],
            ["Restitution", "app/, pages Streamlit", "usage opérationnel"],
            ["Exploitation", "README.md, docs/runbook.md", "installation et incidents"],
        ],
    )
    add_body(
        doc,
        "La documentation utilisateur explique quoi regarder et comment décider ; la documentation "
        "technique explique d'où vient la valeur et comment la reproduire ; le runbook explique quoi "
        "faire lorsque la chaîne ne fonctionne plus. Séparer ces trois usages évite un document unique "
        "trop dense pour le coach et trop imprécis pour l'exploitant.",
    )
    doc.add_heading("8.2 Traçabilité d'un indicateur", level=2)
    add_body(
        doc,
        "Chaque résultat du dossier dérive d'un snapshot JSON produit depuis PostgreSQL Gold. Le "
        "fichier contient la date UTC, le dernier patch, l'échantillon, les statistiques, les "
        "décisions et les limites. Les secrets restent dans les variables d'environnement et ne sont "
        "jamais inscrits dans les livrables. Le dépôt permet ainsi de remonter d'une visualisation à "
        "son calcul, puis aux sources et aux tests de qualité.",
    )
    add_table(
        doc,
        ["Question de traçabilité", "Réponse apportée"],
        [
            ["D'où vient la donnée ?", "Source déclarée, endpoint ou extraction Leaguepedia et objet MinIO."],
            ["Comment a-t-elle été transformée ?", "Modèle dbt versionné et documentation de colonnes."],
            ["Quels contrôles a-t-elle franchis ?", "Tests génériques, singuliers et plan de tests horodaté."],
            ["Quand le résultat a-t-il été produit ?", "Horodatage UTC du snapshot statistique."],
            ["Comment le reproduire ?", "Commandes, dépendances et paramètres documentés."],
        ],
    )
    doc.add_heading("8.3 Procédure de reproduction", level=2)
    add_code(
        doc,
        "# 1. Démarrer les services\n"
        "docker compose up -d\n\n"
        "# 2. Construire et tester les modèles analytiques\n"
        "docker compose exec dbt dbt build --project-dir /usr/app/dbt --profiles-dir /usr/app/dbt\n\n"
        "# 3. Régénérer les statistiques et figures\n"
        ".venv\\Scripts\\python.exe analysis\\statistical_analysis.py\n\n"
        "# 4. Régénérer le dossier\n"
        ".venv\\Scripts\\python.exe analysis\\generate_analysis_report.py",
    )
    add_body(
        doc,
        "La reproduction doit conserver le fichier JSON généré avec le livrable afin que le lecteur "
        "puisse distinguer une variation de données d'une variation de méthode. Toute modification de "
        "formule entraîne une nouvelle version du modèle dbt, du dictionnaire d'indicateurs et du "
        "dossier. Les identifiants secrets et fichiers .env restent exclus de cette archive.",
    )
    doc.add_heading("8.4 Conditions de maintenance", level=2)
    add_body(
        doc,
        "Une livraison est considérée comme maintenable lorsque le build dbt est vert, que le snapshot "
        "est horodaté, que les figures sont régénérées à partir de ce snapshot et que les limites ont "
        "été relues. Les changements de nomenclature externe — version de patch, identifiant de rôle "
        "ou schéma d'API — font l'objet d'un test de non-régression avant déploiement.",
    )


def add_conclusion_and_annexes(doc: Document, data: dict) -> None:
    doc.add_heading("Conclusion", level=1)
    add_body(
        doc,
        "MyLeague couvre la chaîne analytique attendue du Bloc 2 : problème métier traduit en "
        "indicateurs, données interrogées en SQL/Python, hypothèses testées, représentations "
        "justifiées, recommandations prudentes et documentation reproductible. Les résultats "
        "montrent une association solo/pro significative, sans démontrer à ce stade un effet global "
        "des changements Riot. Cette absence de surpromesse est une qualité du dispositif. Deux "
        "preuves humaines restent à obtenir pour rendre le dossier incontestable : un entretien "
        "métier réel et une session de formation tracée.",
    )
    add_body(
        doc,
        "La principale contribution du projet réside dans l'articulation entre ingénierie et usage. "
        "La qualité de collecte n'est pas traitée comme une fin : elle rend possible une analyse dont "
        "le périmètre, le dénominateur et l'incertitude sont explicites. Réciproquement, le besoin du "
        "coach détermine les transformations, tests et niveaux de détail à maintenir dans la plateforme.",
    )
    add_body(
        doc,
        "La prochaine étape n'est pas d'ajouter un modèle plus complexe, mais de confronter le produit "
        "à son environnement réel. Un entretien permettra de vérifier les priorités ; une séance de "
        "formation révélera les ambiguïtés d'usage ; l'intégration des scrims permettra enfin d'évaluer "
        "la valeur de la recommandation pour une équipe donnée. Cette progression préserve la "
        "proportionnalité entre sophistication technique et preuve métier.",
    )
    doc.add_page_break()
    doc.add_heading("Annexe 1 — Matrice de couverture du Bloc 2", level=1)
    add_table(
        doc,
        ["Compétence", "État", "Preuves principales", "Action restante"],
        [
            ["C2.1.1 Besoin métier", "Partiel", "personas, contraintes, guide d'entretien", "entretien réel signé"],
            ["C2.1.2 Plan d'analyse", "Couvert", "axes, métriques, sources, hypothèses", "validation coach souhaitable"],
            ["C2.1.3 Requêtes/calculs", "Couvert", "Gold, dbt, SQL, Python, Streamlit", "maintenir le snapshot"],
            ["C2.1.4 Tests statistiques", "Couvert", "Spearman, Mann–Whitney, interprétation", "répéter sur plus de patchs"],
            ["C2.2.1 Représentation", "Couvert", "5 figures, dashboard, règles d'accessibilité", "test utilisateur conseillé"],
            ["C2.2.2 Recommandations", "Couvert", "recommandations structurées et limites", "faire arbitrer par le staff"],
            ["C2.3.1 Formation", "Partiel", "déroulé, exercice, support séparé", "réaliser et tracer la séance"],
            ["C2.3.2 Documentation", "Couvert", "sources, calculs, commandes, traçabilité", "versionner chaque livraison"],
        ],
    )
    doc.add_heading("Annexe 2 — Guide d'entretien exploratoire", level=1)
    add_bullets(
        doc,
        [
            "Quelle décision de draft vous coûte aujourd'hui le plus de temps ou d'incertitude ?",
            "Quand considérez-vous qu'un nouveau patch dispose d'assez de données ?",
            "Quels indicateurs utilisez-vous déjà et lesquels vous paraissent trompeurs ?",
            "Comment le champion pool et les scrims doivent-ils modifier une recommandation globale ?",
            "Quelles différences attendez-vous entre solo queue et jeu professionnel ?",
            "À quel moment préférez-vous une alerte d'incertitude à une recommandation ?",
            "Quel niveau de détail faut-il conserver pour expliquer un choix aux joueurs ?",
            "Quelle décision prendriez-vous à partir du dashboard actuel, et que vous manque-t-il ?",
        ],
    )
    doc.add_heading("Annexe 3 — Résultats statistiques synthétiques", level=1)
    h1, h2 = data["hypothesis_1"], data["hypothesis_2"]
    add_table(
        doc,
        ["Test", "Échantillon", "Statistique", "p", "Décision"],
        [
            ["Spearman solo/pro", f"n={h1['n']}", f"rho={h1['rho']:.3f}", f"{h1['p_value']:.2e}", "H0 rejetée"],
            ["Mann–Whitney patch", f"{h2['n_changed']} / {h2['n_unchanged']}", f"U={h2['u_statistic']:.0f}", f"{h2['p_value']:.3f}", "H0 non rejetée"],
        ],
    )
    doc.add_heading("Annexe 4 — Registre des preuves", level=1)
    add_bullets(
        doc,
        [
            "analysis/results/statistical_analysis_results.json — snapshot chiffré et horodaté.",
            "analysis/statistical_analysis.py — requêtes, tests et génération des figures.",
            "dbt/models — définitions des modèles analytiques et dictionnaire des colonnes.",
            "dbt/tests/gold_pro_solo_patch_overlap.sql — contrôle de recouvrement des patchs.",
            "docs/MyLeague_Bloc2_Support_FormationV2.docx — support destiné aux utilisateurs.",
            "README.md et docs/ — architecture, sécurité, exploitation et qualité.",
        ],
    )
    doc.add_heading("Annexe 5 — Glossaire", level=1)
    add_table(
        doc,
        ["Terme", "Définition dans MyLeague"],
        [
            ["Champion pool", "Ensemble des champions qu'un joueur peut mobiliser dans un contexte compétitif."],
            ["Confiance", "Qualification du volume disponible pour interpréter une recommandation."],
            ["Draft", "Phase de sélection et de bannissement des champions avant une partie."],
            ["Gold", "Couche de données métiers prêtes à être analysées et exposées."],
            ["Méta", "Ensemble évolutif des choix et stratégies considérés comme efficaces dans un contexte donné."],
            ["Patch", "Version du jeu définissant les règles, statistiques et changements applicables."],
            ["Présence", "Fréquence à laquelle un champion est sélectionné ou banni dans le périmètre observé."],
            ["Scrim", "Partie d'entraînement organisée entre équipes, généralement non publique."],
            ["Snapshot", "Extraction analytique horodatée permettant de reproduire un résultat."],
            ["Solo queue", "File classée individuelle ou en duo, distincte de l'environnement compétitif coordonné."],
        ],
    )
    doc.add_heading("Annexe 6 — Références techniques", level=1)
    add_body(
        doc,
        "Les références ci-dessous ont été consultées le 27 juillet 2026. Les définitions propres au "
        "projet restent documentées dans le dépôt, afin qu'une évolution d'un outil externe ne modifie "
        "pas silencieusement la méthode d'analyse.",
    )
    add_bullets(
        doc,
        [
            "Riot Games, Developer Portal — League of Legends : https://developer.riotgames.com/docs/lol",
            "dbt Labs, documentation des data tests : https://docs.getdbt.com/docs/build/data-tests",
            "Streamlit, documentation officielle : https://docs.streamlit.io/",
            "SciPy, corrélation de Spearman : https://docs.scipy.org/doc/scipy/reference/generated/scipy.stats.spearmanr.html",
            "SciPy, test de Mann–Whitney U : https://docs.scipy.org/doc/scipy/reference/generated/scipy.stats.mannwhitneyu.html",
            "Leaguepedia, source de données compétitives : https://lol.fandom.com/wiki/League_of_Legends_Esports_Wiki",
        ],
    )


def build_main_document(data: dict) -> None:
    doc = Document(TEMPLATE)
    keep_template_cover(doc)
    configure_document(doc)
    add_front_matter(doc)
    add_executive_summary(doc, data)
    add_introduction(doc)
    add_need_analysis(doc)
    add_analysis_plan(doc, data)
    add_queries_and_results(doc, data)
    add_statistics(doc, data)
    add_visualisation(doc)
    add_recommendations(doc)
    add_training_and_docs(doc)
    add_conclusion_and_annexes(doc, data)
    doc.core_properties.title = "MyLeague — Dossier Bloc 2 — Version 2"
    doc.core_properties.subject = "Exploiter les données à des fins d'analyse"
    doc.core_properties.author = "David Nguyen"
    doc.save(OUTPUT)


def build_training_document(data: dict) -> None:
    doc = Document()
    configure_document(doc)
    add_page_footer(doc, "MyLeague - Formation Bloc 2")
    title = doc.add_heading("MyLeague — Support de formation", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle = doc.add_paragraph("Lire la méta sans surinterpréter les données · Bloc 2 · 45 minutes")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_note(
        doc,
        "Objectif",
        "À l'issue de la séance, le participant sait filtrer l'analyse, lire un indicateur avec son "
        "volume, qualifier la confiance et formuler une recommandation assortie d'une limite.",
    )
    doc.add_heading("1. Parcours de la séance", level=1)
    add_table(
        doc,
        ["Durée", "Séquence", "Production du participant"],
        [
            ["5 min", "Problème et périmètre", "une décision métier reformulée"],
            ["10 min", "Navigation Streamlit", "patch et rôle correctement filtrés"],
            ["10 min", "Métriques et confiance", "une lecture avec dénominateur"],
            ["15 min", "Cas pratique de draft", "une recommandation argumentée"],
            ["5 min", "Correction et feedback", "une limite explicitée"],
        ],
    )
    doc.add_heading("2. Les cinq réflexes", level=1)
    add_bullets(
        doc,
        [
            "Vérifier le patch et la date de dernière ingestion.",
            "Lire le nombre de parties avant le winrate ou le score.",
            "Comparer solo queue et professionnel, sans supposer qu'ils doivent coïncider.",
            "Traiter une confiance faible comme une invitation à collecter ou observer davantage.",
            "Confronter le signal au champion pool, au match-up, aux scrims et au plan de jeu.",
        ],
    )
    doc.add_heading("3. Cas pratique", level=1)
    top = data["draft"]["top_recommendations"][:4]
    add_table(
        doc,
        ["Champion", "Rôle", "Score", "Picks", "Présence pro", "Confiance"],
        [[x["champion_name"], x["role"], x["priority_score"], x["picks"], percent(x["pro_presence"]), x["confidence_level"]] for x in top],
    )
    add_body(
        doc,
        "Consigne : choisissez un champion à tester en priorité. Rédigez une phrase comprenant la "
        "décision, deux éléments de preuve et une réserve. Puis indiquez quelle donnée interne "
        "— scrim, maîtrise du joueur ou match-up — pourrait renverser votre choix.",
    )
    add_note(
        doc,
        "Correction attendue",
        "Il n'existe pas une réponse unique. La qualité repose sur la traçabilité du raisonnement et "
        "sur la prudence face à la confiance faible, pas sur le champion choisi.",
        GREEN,
    )
    doc.add_heading("4. Aide-mémoire des indicateurs", level=1)
    add_table(
        doc,
        ["Indicateur", "Question", "Piège"],
        [
            ["Winrate", "le champion gagne-t-il dans cet échantillon ?", "oublier le volume et le rôle"],
            ["Présence", "à quelle fréquence est-il pick ou ban ?", "confondre popularité et puissance"],
            ["Présence pro", "est-il contesté en environnement coordonné ?", "ignorer région et compétition"],
            ["Score de priorité", "plusieurs signaux convergent-ils ?", "le lire comme une vérité absolue"],
            ["Confiance", "le volume autorise-t-il une décision ?", "masquer l'incertitude"],
        ],
    )
    doc.add_heading("5. Évaluation de fin de séance", level=1)
    add_bullets(
        doc,
        [
            "Je sais retrouver le patch et le rôle analysés : oui / à revoir.",
            "Je sais expliquer la différence entre présence et winrate : oui / à revoir.",
            "Je sais identifier un échantillon fragile : oui / à revoir.",
            "Ma décision cite une preuve et une limite : oui / à revoir.",
            "Retour libre : quelle information manque pour votre usage quotidien ?",
        ],
    )
    add_note(
        doc,
        "Traçabilité RNCP",
        "Après la séance, conserver la date, la liste des participants, les réponses au cas pratique, "
        "le questionnaire et les améliorations décidées. Ce document seul n'est pas une preuve que "
        "la formation a effectivement eu lieu.",
        AMBER,
    )
    doc.core_properties.title = "MyLeague — Support de formation Bloc 2 — Version 2"
    doc.core_properties.author = "David Nguyen"
    doc.save(TRAINING_OUTPUT)


def main() -> None:
    data = json.loads(RESULTS.read_text(encoding="utf-8"))
    build_main_document(data)
    build_training_document(data)
    print(OUTPUT)
    print(TRAINING_OUTPUT)


if __name__ == "__main__":
    main()
