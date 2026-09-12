from pathlib import Path

from docx import Document
from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / "docs"


def dump_docx(path: Path) -> None:
    doc = Document(path)
    out = [f"DOCUMENT: {path.name}", f"paragraphs={len(doc.paragraphs)} tables={len(doc.tables)}"]
    for i, paragraph in enumerate(doc.paragraphs):
        text = paragraph.text.strip()
        if text:
            style = paragraph.style.name if paragraph.style else ""
            out.append(f"P{i:04d} [{style}] {text}")
    for ti, table in enumerate(doc.tables):
        out.append(f"TABLE {ti} rows={len(table.rows)} cols={len(table.columns)}")
        for ri, row in enumerate(table.rows):
            cells = [" ".join(cell.text.split()) for cell in row.cells]
            out.append(f"T{ti:02d}R{ri:03d} | " + " | ".join(cells))
    target = ROOT / ".cache" / f"{path.stem}_dump.txt"
    target.parent.mkdir(exist_ok=True)
    target.write_text("\n".join(out), encoding="utf-8")
    print(target)


def dump_xlsx(path: Path) -> None:
    wb = load_workbook(path, data_only=True, read_only=True)
    out = [f"WORKBOOK: {path}", f"Sheets: {wb.sheetnames}"]
    for ws in wb.worksheets:
        out.append(f"SHEET {ws.title} {ws.max_row}x{ws.max_column}")
        for row in ws.iter_rows():
            values = [str(cell.value).strip() if cell.value is not None else "" for cell in row]
            if any(values):
                out.append(" | ".join(values))
    target = ROOT / ".cache" / "evaluation_grid_dump.txt"
    target.write_text("\n".join(out), encoding="utf-8")
    print(target)


if __name__ == "__main__":
    for name in (
        "Bloc3_David_Nguyen_M2.docx",
        "Bloc2_David_Nguyen_M2_consolide.docx",
        "Bloc2_David_Nguyen_M2_corrige.docx",
    ):
        dump_docx(DOCS / name)
    dump_xlsx(Path(r"C:\Users\nguye\Downloads\Vrai grille d'eval pfe(1).xlsx"))
