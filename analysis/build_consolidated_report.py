"""Construit le dossier Bloc 2 consolidé avec les preuves multi-patch."""

from __future__ import annotations

import json
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches

import analysis.update_competency_evidence as base


ROOT = Path(__file__).resolve().parents[1]
MULTI_RESULTS = ROOT / "analysis" / "results" / "multi_patch_analysis_results.json"
ASSETS = ROOT / "docs" / "assets" / "bloc2"
OUTPUT = ROOT / "docs" / "Bloc2_David_Nguyen_M2_consolide.docx"


def fill_row(row, values) -> None:
    for cell, value in zip(row.cells, values, strict=True):
        base.replace_cell_text(cell, str(value))


def add_centered_figure(doc: Document, filename: str, caption: str) -> None:
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.add_run().add_picture(str(ASSETS / filename), width=Inches(6.2))
    caption_paragraph = doc.add_paragraph(caption)
    caption_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER


def main() -> None:
    # Applique d'abord toutes les corrections du corps du dossier.
    base.main()
    doc = Document(base.OUTPUT)
    current = json.loads(base.RESULTS.read_text(encoding="utf-8"))
    multi = json.loads(MULTI_RESULTS.read_text(encoding="utf-8"))
    patches = current["patches"]
    previous_patch = patches[-2]["patch"]
    current_patch = patches[-1]["patch"]
    h2 = current["hypothesis_2"]

    base.replace_paragraph_starting_with(
        doc,
        "La photographie calculée couvre",
        "La photographie calculée couvre "
        + ", ".join(
            f"{item['total_matches']} parties au patch {item['patch']}" for item in patches
        )
        + f". Sur le patch {current['latest_patch']}, "
        f"{current['draft']['champions_analyzed']} couples champion-rôle ayant au moins "
        "10 picks sont analysés. Les résultats sont enregistrés dans un snapshot JSON "
        f"horodaté {current['generated_at_utc']} et ne sont ni estimés ni laissés à compléter.",
    )
    base.replace_paragraph_starting_with(
        doc,
        "Le test de Mann–Whitney bilatéral compare",
        f"Le test de Mann–Whitney bilatéral compare l'amplitude absolue des variations "
        f"entre les patchs {previous_patch} et {current_patch} pour "
        f"{h2['n_changed']} champions modifiés et {h2['n_unchanged']} champions témoins. "
        f"Les amplitudes moyennes sont respectivement de "
        f"{h2['changed_mean_absolute_delta'] * 100:.1f} et "
        f"{h2['unchanged_mean_absolute_delta'] * 100:.1f} points de pourcentage ; "
        f"U = {h2['u_statistic']:.0f} et p = {h2['p_value']:.3f}. H0 n'est pas "
        "rejetée au seuil de 5 % : une différence n'est pas démontrée. Le résultat "
        "est interprété comme une incertitude et non comme une preuve d'absence d'effet.",
    )
    base.replace_paragraph_starting_with(
        doc,
        "Figure 4 : Comparaison calculée",
        f"Figure 4 : comparaison calculée des variations de winrate entre les patchs "
        f"{previous_patch} et {current_patch} "
        f"(U = {h2['u_statistic']:.0f}, p = {h2['p_value']:.3f}).",
    )

    doc.add_page_break()
    doc.add_heading("Annexe 7 — Consolidation multi-patch", level=1)
    doc.add_paragraph(
        "Cette annexe étend l'analyse principale à toutes les fenêtres de patch "
        "disponibles au moment de l'extraction. Les résultats proviennent des mêmes "
        "modèles Gold que les dashboards et restent exploratoires."
    )

    doc.add_heading("A7.1 Couverture des données", level=2)
    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    fill_row(
        table.rows[0],
        ["Patch", "Parties solo queue EUW", "Parties professionnelles"],
    )
    for item in multi["coverage"]:
        fill_row(
            table.add_row(),
            [item["patch"], item["solo_matches"], item["pro_matches"]],
        )
    add_centered_figure(
        doc,
        "multi_patch_coverage.png",
        "Figure 6 : couverture solo queue et professionnelle par patch.",
    )

    doc.add_heading("A7.2 Stabilité de l'association solo/pro", level=2)
    table = doc.add_table(rows=1, cols=5)
    table.style = "Table Grid"
    fill_row(table.rows[0], ["Patch", "n", "ρ", "p", "Décision"])
    for item in multi["spearman_by_patch"]:
        p_text = (
            "< 0,001"
            if item.get("p_value", 1) < 0.001
            else f"{item['p_value']:.3f}".replace(".", ",")
        )
        fill_row(
            table.add_row(),
            [
                item["patch"],
                item["n"],
                f"{item['rho']:.3f}".replace(".", ","),
                p_text,
                item["decision"],
            ],
        )
    doc.add_paragraph(
        "H0 est rejetée sur chacun des trois patchs communs : l'association entre "
        "présence solo queue et présence professionnelle est positive et retrouvée "
        "dans plusieurs fenêtres temporelles. Son intensité varie selon le patch ; "
        "les deux indicateurs restent donc présentés séparément."
    )
    add_centered_figure(
        doc,
        "spearman_by_patch.png",
        "Figure 7 : corrélation de Spearman solo/pro calculée par patch.",
    )

    pooled = multi["mann_whitney_pooled_transitions"]
    doc.add_heading("A7.3 Effet regroupé des changements officiels", level=2)
    doc.add_paragraph(
        f"Le test regroupé porte sur {pooled['n_changed']} observations champion-patch "
        f"modifiées et {pooled['n_control']} observations témoins. La médiane de "
        f"|Δ winrate| vaut {pooled['median_changed'] * 100:.1f} points chez les "
        f"champions modifiés contre {pooled['median_control'] * 100:.1f} points chez "
        f"les témoins. Le test produit U = {pooled['u_statistic']:.0f}, "
        f"p = {pooled['p_value']:.3f} et une corrélation rank-bisériale de "
        f"{pooled['rank_biserial']:.3f}. H0 n'est pas rejetée. L'intervalle bootstrap "
        "à 95 % de la différence de médianes "
        f"[{pooled['median_difference_bootstrap_ci95'][0] * 100:.1f} ; "
        f"{pooled['median_difference_bootstrap_ci95'][1] * 100:.1f}] contient zéro : "
        "l'effet des changements officiels n'est pas démontré avec cet échantillon."
    )
    add_centered_figure(
        doc,
        "multi_patch_change_effect.png",
        "Figure 8 : variations avec et sans changement officiel, toutes transitions.",
    )

    doc.add_heading("A7.4 Traçabilité et limites", level=2)
    for item in multi["limitations"]:
        doc.add_paragraph(item, style="List Bullet")
    doc.add_paragraph(
        "Preuves reproductibles : jobs/riot/historical_backfill.py, "
        "analysis/multi_patch_analysis.py, "
        "analysis/results/multi_patch_analysis_results.json et "
        "dbt/target/run_results.json."
    )

    doc.core_properties.title = "MyLeague — Dossier Bloc 2 consolidé"
    doc.core_properties.subject = (
        "Requêtes, dashboards, résultats et tests statistiques multi-patch"
    )
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
