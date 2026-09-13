"""Contrôles automatisés du DOCX et de son rendu PDF."""

from pathlib import Path
from zipfile import ZipFile

from docx import Document
from pypdf import PdfReader


ROOT = Path(__file__).resolve().parents[1]
DOCX = ROOT / "docs" / "Bloc3_David_Nguyen_M2_final.docx"
PDF = ROOT / ".cache" / "bloc3_render" / "Bloc3_David_Nguyen_M2_final.pdf"


def main() -> None:
    doc = Document(DOCX)
    texts = [paragraph.text for paragraph in doc.paragraphs]
    texts.extend(
        cell.text
        for table in doc.tables
        for row in table.rows
        for cell in row.cells
    )
    xml = ZipFile(DOCX).read("word/document.xml")
    combined = "\n".join(texts)
    body_paragraphs = []
    inside_body = False
    for paragraph in doc.paragraphs:
        if paragraph.text == "Introduction":
            inside_body = True
        if paragraph.text == "Annexes":
            inside_body = False
        if (
            inside_body
            and paragraph.text.strip()
            and paragraph.style
            and paragraph.style.name == "normal"
        ):
            body_paragraphs.append(paragraph)

    reader = PdfReader(PDF)
    pdf_texts = [page.extract_text() or "" for page in reader.pages]
    annex_page = next(
        page
        for page, text in enumerate(pdf_texts, start=1)
        if text.strip().endswith("Annexes")
    )

    checks = {
        "DOCX lisible": len(doc.paragraphs) > 0,
        "23 tableaux conservés": len(doc.tables) == 23,
        "aucun point-virgule visible": not any(";" in text for text in texts),
        "aucun double tiret visible": not any("--" in text for text in texts),
        "aucun numéro provisoire": not any(
            "À actualiser" in text or "\tXX" in text for text in texts
        ),
        "aucun Remerciements fantôme": b"Remerciements" not in xml,
        "ancien champ TOC supprimé": b"TOC" not in xml,
        "lignes de tableaux non fractionnables": xml.count(b"cantSplit")
        == sum(len(table.rows) for table in doc.tables),
        "en-têtes de tableaux répétables": xml.count(b"tblHeader") >= len(doc.tables),
        "20 pages maximum hors garde et annexes": annex_page - 2 <= 20,
        "36 pages totales au rendu": len(reader.pages) == 36,
        "snapshot canonique du 28 juillet": "28 juillet 2026 à 00 h 01 UTC" in combined,
        "tests statistiques canoniques": all(
            token in combined
            for token in ("137 champions", "6 champions modifiés", "122 témoins", "U = 276", "p = 0,313")
        ),
        "preuve dbt canonique": all(
            token in combined
            for token in ("18 modèles", "63 tests de données", "81 opérations réussies")
        ),
        "tests utilisateurs non surévalués": "test avec le staff Nexus non encore exécuté" in combined,
        "formation non surévaluée": "Session Nexus non encore réalisée" in combined,
        "chronologie S5 et S6": all(
            token in combined
            for token in ("arbitrage ouvert le 22 juillet", "validation le 28 juillet en S6")
        ),
        "libellé de charge non ambigu": "52 jours pour la phase intensive et 70 jours pour le projet complet" in combined,
        "calcul 21, 18 et 13 jours": "Répartition calculée de 21, 18 et 13 jours" in combined,
        "interligne 1,5 dans le corps": all(
            paragraph.paragraph_format.line_spacing == 1.5
            for paragraph in body_paragraphs
        ),
        "alinéas dans le corps": all(
            paragraph.paragraph_format.first_line_indent is not None
            and paragraph.paragraph_format.first_line_indent.cm >= 0.69
            for paragraph in body_paragraphs
        ),
    }
    for label, passed in checks.items():
        print(f"{'OK' if passed else 'ECHEC'} | {label}")
    print(f"Pages du corps hors garde: {annex_page - 2}")
    print(f"Première page des annexes: {annex_page}")
    print(f"Marqueurs d'en-tête de tableau: {xml.count(b'tblHeader')}")
    if not all(checks.values()):
        raise SystemExit(1)


if __name__ == "__main__":
    main()
